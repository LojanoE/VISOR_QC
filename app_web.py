import json
import os
import sys
import threading
import time
import uuid
from datetime import datetime, timedelta
from collections import defaultdict
from io import BytesIO

from flask import (Flask, jsonify, request, send_file, send_from_directory,
                   Response, url_for, session, redirect)
from werkzeug.security import safe_join

import database
import utils
import pyproj
from PIL import Image

app = Flask(__name__)
app.secret_key = uuid.uuid4().hex

cors_origins = os.environ.get('CORS_ORIGINS', '')
if cors_origins:
    from flask_cors import CORS
    CORS(app, origins=[o.strip() for o in cors_origins.split(',') if o.strip()], supports_credentials=True)

if os.environ.get('CROSS_ORIGIN', '').lower() == 'true':
    app.config['SESSION_COOKIE_SAMESITE'] = 'None'
    app.config['SESSION_COOKIE_SECURE'] = True

transformer = pyproj.Transformer.from_crs("epsg:4326", "epsg:24877", always_xy=True)

scan_state = {"running": False, "progress": 0, "message": "", "total": 0, "current": 0}


# ──────────────── Authentication ────────────────

@app.before_request
def check_login():
    allowed_routes = ['login_page', 'login_html', 'index_page', 'api_login', 'api_session', 'api_logout', 'static']
    if request.endpoint and request.endpoint not in allowed_routes and 'static' not in (request.endpoint or ''):
        if 'user' not in session:
            if request.path.startswith('/api/'):
                return jsonify({'error': 'No autenticado'}), 401
            return redirect(url_for('login_html'))


@app.route('/login', methods=['GET', 'POST'])
def login_page():
    return redirect(url_for('login_html'))


@app.route('/api/login', methods=['POST'])
def api_login():
    data = request.get_json(silent=True) or {}
    username = data.get('username', '').strip()
    password = data.get('password', '')
    try:
        database.init_db()
        database.init_users_table()
        if database.verify_user(username, password):
            session['user'] = username
            session.permanent = True
            return jsonify({'success': True, 'username': username})
        else:
            return jsonify({'error': 'Usuario o contraseña incorrectos'}), 401
    except Exception as e:
        return jsonify({'error': f'Error de conexión: {e}'}), 500


@app.route('/api/session')
def api_session():
    if 'user' in session:
        return jsonify({'authenticated': True, 'username': session['user']})
    return jsonify({'authenticated': False}), 401


@app.route('/api/logout')
def api_logout():
    session.pop('user', None)
    return jsonify({'success': True})


def load_config():
    config_path = database.get_app_data_path('config.json')
    if os.path.exists(config_path):
        with open(config_path, 'r', encoding='utf-8') as f:
            try:
                return json.load(f)
            except json.JSONDecodeError:
                return {}
    return {}


def save_config(config):
    config_path = database.get_app_data_path('config.json')
    with open(config_path, 'w', encoding='utf-8') as f:
        json.dump(config, f, indent=4)


def parse_date(date_str):
    if not date_str:
        return None
    try:
        return datetime.strptime(date_str, '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return None


def photo_to_dict(photo):
    return {
        'path': photo.get('path', ''),
        'date': photo['date'].isoformat() if photo.get('date') and isinstance(photo['date'], datetime) else str(photo.get('date', '')),
        'lat': photo['coords'][0] if photo.get('coords') else None,
        'lon': photo['coords'][1] if photo.get('coords') else None,
        'utm_x': photo.get('utm_x'),
        'utm_y': photo.get('utm_y'),
        'work_front': photo.get('work_front', ''),
        'coronation': photo.get('coronation', ''),
        'activity_performed': photo.get('activity_performed', ''),
        'observation_category': photo.get('observation_category', ''),
    }


def normalize_metadata_value(val, field_title):
    if val and isinstance(val, str) and val.strip():
        return val.strip()
    return f"Sin {field_title}"


# ──────────────────────────────── Routes ────────────────────────────────

@app.route('/')
def index_page():
    return send_from_directory(app.root_path, 'index.html')


@app.route('/login.html')
def login_html():
    return send_from_directory(app.root_path, 'login.html')


# ── Config ──

@app.route('/api/config', methods=['GET'])
def get_config():
    config = load_config()
    return jsonify({
        'map_image_path': config.get('map_image_path', ''),
        'coords_file_path': config.get('coords_file_path', ''),
        'photos_folder_path': config.get('photos_folder_path', ''),
        'docs_folder_path': config.get('docs_folder_path', ''),
    })


@app.route('/api/config', methods=['POST'])
def set_config():
    data = request.get_json()
    config = load_config()
    for key in ('map_image_path', 'coords_file_path', 'photos_folder_path', 'docs_folder_path'):
        if key in data:
            config[key] = data[key]
    save_config(config)
    return jsonify({'status': 'ok'})


@app.route('/api/config/paths', methods=['GET'])
def browse_paths():
    base = request.args.get('dir', '')
    if not base:
        drives = []
        if os.name == 'nt':
            import string
            for letter in string.ascii_uppercase:
                p = f'{letter}:\\'
                if os.path.exists(p):
                    drives.append(p)
        else:
            drives = ['/']
        return jsonify({'dirs': drives, 'parent': ''})
    base = base.replace('/', os.sep)
    if not os.path.isdir(base):
        return jsonify({'dirs': [], 'parent': ''})
    try:
        entries = []
        for e in os.listdir(base):
            full = os.path.join(base, e)
            if os.path.isdir(full):
                try:
                    entries.append(full)
                except PermissionError:
                    pass
    except PermissionError:
        entries = []
    parent = os.path.dirname(base)
    return jsonify({'dirs': sorted(entries), 'parent': parent if parent != base else ''})


# ── Map ──

@app.route('/api/map/bounds')
def map_bounds():
    config = load_config()
    coords_path = config.get('coords_file_path', '')
    if not coords_path or not os.path.exists(coords_path):
        return jsonify({'error': 'Archivo de coordenadas no configurado'}), 404
    bounds = utils.parse_coords_file(coords_path)
    if not bounds:
        return jsonify({'error': 'No se pudo leer el archivo de coordenadas'}), 400
    return jsonify(bounds)


@app.route('/api/map/image')
def map_image():
    config = load_config()
    map_path = config.get('map_image_path', '')
    if not map_path or not os.path.exists(map_path):
        return jsonify({'error': 'Imagen del mapa no configurada'}), 404
    return send_file(map_path, mimetype='image/jpeg')


# ── Photos ──

@app.route('/api/photos')
def get_photos():
    config = load_config()
    photos_folder = config.get('photos_folder_path', '')

    if not database.db_exists():
        return jsonify([])

    try:
        database.init_db()
    except Exception:
        pass

    all_photos = database.get_photos_for_indexing_with_filters(has_coords=True)

    date_from_str = request.args.get('date_from')
    date_to_str = request.args.get('date_to')
    shift = request.args.get('shift', 'Ambos')
    work_front = request.args.getlist('work_front')
    coronation = request.args.getlist('coronation')
    activity_performed = request.args.getlist('activity_performed')
    observation_category = request.args.getlist('observation_category')

    date_from = parse_date(date_from_str)
    date_to = parse_date(date_to_str)

    field_titles = {
        'work_front': 'Frente de Trabajo',
        'coronation': 'Coronamiento',
        'activity_performed': 'Actividad Realizada',
        'observation_category': 'Categoría de Observación',
    }

    filter_selections = {
        'work_front': set(work_front) if work_front else set(),
        'coronation': set(coronation) if coronation else set(),
        'activity_performed': set(activity_performed) if activity_performed else set(),
        'observation_category': set(observation_category) if observation_category else set(),
    }

    filtered = []
    for photo in all_photos:
        if not photo.get('coords'):
            continue

        if photos_folder and not photo['path'].startswith(photos_folder):
            continue

        photo_date = photo.get('date')
        if not photo_date:
            continue
        if isinstance(photo_date, str):
            try:
                photo_date = datetime.fromisoformat(photo_date)
            except (ValueError, TypeError):
                continue

        if date_from and date_to:
            if photo_date.date() < date_from or photo_date.date() > date_to:
                continue
            hour = photo_date.hour
            if shift == 'Diurno' and not (7 <= hour < 19):
                continue
            if shift == 'Nocturno' and not (hour >= 19 or hour < 7):
                continue
        elif shift == 'Diurno':
            hour = photo_date.hour
            if not (7 <= hour < 19):
                continue
        elif shift == 'Nocturno':
            hour = photo_date.hour
            if not (hour >= 19 or hour < 7):
                continue

        skip = False
        for field, selected in filter_selections.items():
            if not selected:
                continue
            val = photo.get(field)
            title = field_titles[field]
            norm_val = normalize_metadata_value(val, title)
            if norm_val not in selected:
                skip = True
                break
        if skip:
            continue

        if not photo.get('utm_x') or not photo.get('utm_y'):
            lat, lon = photo['coords']
            try:
                utm_x, utm_y = transformer.transform(lon, lat)
                photo['utm_x'] = utm_x
                photo['utm_y'] = utm_y
            except Exception:
                continue

        filtered.append(photo)

    return jsonify([photo_to_dict(p) for p in filtered])


@app.route('/api/filters/options')
def filter_options():
    config = load_config()
    photos_folder = config.get('photos_folder_path', '')

    empty = {f: {'values': [], 'counts': {}} for f in FILTER_FIELDS_LIST}
    if not database.db_exists():
        return jsonify(empty)

    try:
        database.init_db()
    except Exception:
        return jsonify(empty)

    all_photos = database.get_photos_for_indexing_with_filters(has_coords=True)

    date_from_str = request.args.get('date_from')
    date_to_str = request.args.get('date_to')
    shift = request.args.get('shift', 'Ambos')

    date_from = parse_date(date_from_str)
    date_to = parse_date(date_to_str)

    selection_params = {
        'work_front': set(request.args.getlist('work_front')),
        'coronation': set(request.args.getlist('coronation')),
        'activity_performed': set(request.args.getlist('activity_performed')),
        'observation_category': set(request.args.getlist('observation_category')),
    }

    field_titles = {
        'work_front': 'Frente de Trabajo',
        'coronation': 'Coronamiento',
        'activity_performed': 'Actividad Realizada',
        'observation_category': 'Categoría de Observación',
    }

    def passes_date_shift(photo, df, dt, sh):
        pd = photo.get('date')
        if not pd:
            return False
        if isinstance(pd, str):
            try:
                pd = datetime.fromisoformat(pd)
            except (ValueError, TypeError):
                return False
        if df and dt:
            if pd.date() < df or pd.date() > dt:
                return False
            h = pd.hour
            if sh == 'Diurno' and not (7 <= h < 19):
                return False
            if sh == 'Nocturno' and not (h >= 19 or h < 7):
                return False
        elif sh == 'Diurno':
            if not (7 <= pd.hour < 19):
                return False
        elif sh == 'Nocturno':
            if not (pd.hour >= 19 or pd.hour < 7):
                return False
        return True

    def passes_metadata_filters(photo, exclude_field, sels, ftitles):
        for f, sel in sels.items():
            if f == exclude_field or not sel:
                continue
            val = photo.get(f)
            norm = normalize_metadata_value(val, ftitles[f])
            if norm not in sel:
                return False
        return True

    result = {}
    for field in FILTER_FIELDS_LIST:
        counts = {}
        for photo in all_photos:
            if not photo.get('coords'):
                continue
            if photos_folder and not photo['path'].startswith(photos_folder):
                continue
            if not passes_date_shift(photo, date_from, date_to, shift):
                continue
            if not passes_metadata_filters(photo, field, selection_params, field_titles):
                continue
            val = photo.get(field)
            norm = normalize_metadata_value(val, field_titles[field])
            counts[norm] = counts.get(norm, 0) + 1

        sorted_vals = sorted(counts.keys())
        result[field] = {'values': sorted_vals, 'counts': counts}

    return jsonify(result)


FILTER_FIELDS_LIST = ['work_front', 'coronation', 'activity_performed', 'observation_category']


@app.route('/api/photos/daterange')
def photo_date_range():
    config = load_config()
    photos_folder = config.get('photos_folder_path', '')

    if not database.db_exists():
        return jsonify({'min_date': None, 'max_date': None})

    try:
        database.init_db()
    except Exception:
        return jsonify({'min_date': None, 'max_date': None})

    all_photos = database.get_photos_for_indexing_with_filters(has_coords=True)

    dates = []
    for photo in all_photos:
        if photos_folder and not photo['path'].startswith(photos_folder):
            continue
        pd = photo.get('date')
        if pd:
            if isinstance(pd, str):
                try:
                    pd = datetime.fromisoformat(pd)
                except (ValueError, TypeError):
                    continue
            dates.append(pd.date())

    if not dates:
        return jsonify({'min_date': None, 'max_date': None})

    return jsonify({
        'min_date': min(dates).isoformat(),
        'max_date': max(dates).isoformat()
    })


@app.route('/api/photo/image')
def photo_image():
    photo_path = request.args.get('path', '')
    if not photo_path or not os.path.exists(photo_path):
        return jsonify({'error': 'Imagen no encontrada'}), 404
    config = load_config()
    photos_folder = config.get('photos_folder_path', '')
    if photos_folder and not photo_path.startswith(photos_folder):
        return jsonify({'error': 'Acceso denegado'}), 403
    return send_file(photo_path, mimetype='image/jpeg')


@app.route('/api/photo/thumbnail')
def photo_thumbnail():
    photo_path = request.args.get('path', '')
    if not photo_path or not os.path.exists(photo_path):
        return jsonify({'error': 'Imagen no encontrada'}), 404
    config = load_config()
    photos_folder = config.get('photos_folder_path', '')
    if photos_folder and not photo_path.startswith(photos_folder):
        return jsonify({'error': 'Acceso denegado'}), 403
    try:
        img = Image.open(photo_path)
        img.thumbnail((150, 150), Image.Resampling.LANCZOS)
        buf = BytesIO()
        img.save(buf, format='JPEG', quality=75)
        buf.seek(0)
        return send_file(buf, mimetype='image/jpeg')
    except Exception:
        return send_file(photo_path, mimetype='image/jpeg')


# ── Metadata ──

@app.route('/api/photo/metadata', methods=['PUT'])
def update_metadata():
    data = request.get_json()
    photo_path = data.get('path', '')
    if not photo_path or not os.path.exists(photo_path):
        return jsonify({'error': 'Imagen no encontrada'}), 404

    new_data = {
        'work_front': data.get('work_front', ''),
        'coronation': data.get('coronation', ''),
        'activity_performed': data.get('activity_performed', ''),
        'observation_category': data.get('observation_category', ''),
    }

    try:
        existing = database.get_photo(photo_path)
        if existing:
            full_meta = existing.copy()
            full_meta.update(new_data)
        else:
            full_meta = new_data.copy()

        modified_time = datetime.fromtimestamp(os.path.getmtime(photo_path))
        database.upsert_photo(photo_path, modified_time, full_meta)
    except Exception as e:
        return jsonify({'error': f'Error al actualizar DB: {e}'}), 500

    try:
        utils.update_image_metadata(photo_path, new_data)
    except Exception:
        pass

    return jsonify({'status': 'ok'})


# ── Cache / Scan ──

@app.route('/api/cache/scan', methods=['POST'])
def start_scan():
    global scan_state
    if scan_state['running']:
        return jsonify({'status': 'already_running'})

    config = load_config()
    photos_folder = config.get('photos_folder_path', '')
    if not photos_folder:
        return jsonify({'error': 'Carpeta de fotos no configurada'}), 400

    scan_state = {"running": True, "progress": 0, "message": "Escaneando...", "total": 0, "current": 0}

    def scan_worker():
        global scan_state
        try:
            database.init_db()
            image_files = list(utils.find_image_files(photos_folder))
            scan_state['total'] = len(image_files)
            scan_state['current'] = 0

            batch = []
            for i, img_path in enumerate(image_files):
                try:
                    metadata = utils.get_image_metadata(img_path)
                    if metadata and metadata.get('coords'):
                        ux, uy = transformer.transform(metadata['coords'][1], metadata['coords'][0])
                        metadata['utm_x'] = ux
                        metadata['utm_y'] = uy
                        file_modified = datetime.fromtimestamp(os.path.getmtime(img_path))
                        batch.append((img_path, file_modified, metadata))
                except Exception:
                    pass

                scan_state['current'] = i + 1
                scan_state['progress'] = int((i + 1) / len(image_files) * 100)

                if len(batch) >= 100:
                    try:
                        database.bulk_upsert_photos(batch)
                    except Exception:
                        pass
                    batch = []

            if batch:
                try:
                    database.bulk_upsert_photos(batch)
                except Exception:
                    pass

            scan_state['message'] = 'Completado'
            scan_state['progress'] = 100
        except Exception as e:
            scan_state['message'] = f'Error: {e}'
        finally:
            scan_state['running'] = False

    t = threading.Thread(target=scan_worker, daemon=True)
    t.start()
    return jsonify({'status': 'started'})


@app.route('/api/cache/progress')
def scan_progress():
    def generate():
        while True:
            data = json.dumps(scan_state)
            yield f"data: {data}\n\n"
            if not scan_state['running']:
                break
            time.sleep(0.5)

    return Response(generate(), mimetype='text/event-stream',
                    headers={'Cache-Control': 'no-cache', 'X-Accel-Buffering': 'no'})


@app.route('/api/cache/status')
def cache_status():
    return jsonify(scan_state)


# ── Documents ──

@app.route('/api/documents')
def get_documents():
    config = load_config()
    docs_path = config.get('docs_folder_path', '')
    if not docs_path or not os.path.exists(docs_path):
        return jsonify([])

    date_from_str = request.args.get('date_from')
    date_to_str = request.args.get('date_to')

    visible_dates = set()
    all_photos = database.get_photos_for_indexing_with_filters(has_coords=True)
    for photo in all_photos:
        d = photo.get('date')
        if d:
            if isinstance(d, str):
                try:
                    d = datetime.fromisoformat(d)
                except (ValueError, TypeError):
                    continue
            visible_dates.add(d.date())

    matching = []
    for root, dirs, files in os.walk(docs_path):
        for f in files:
            if f.startswith('.'):
                continue
            if f.lower().endswith(('.pdf', '.doc', '.docx', '.txt', '.xls', '.xlsx')):
                doc_date = utils.parse_document_name(f)
                if doc_date and doc_date.date() in visible_dates:
                    matching.append({
                        'name': f,
                        'path': os.path.join(root, f),
                    })

    return jsonify(matching)


@app.route('/api/document/open')
def open_document():
    doc_path = request.args.get('path', '')
    if not doc_path or not os.path.exists(doc_path):
        return jsonify({'error': 'Documento no encontrado'}), 404
    return send_file(doc_path)


# ── Export Excel ──

@app.route('/api/export/excel')
def export_excel():
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    config = load_config()

    if not database.db_exists():
        return jsonify({'error': 'Base de datos no disponible'}), 400

    date_from_str = request.args.get('date_from')
    date_to_str = request.args.get('date_to')
    shift = request.args.get('shift', 'Ambos')
    work_front = request.args.getlist('work_front')
    coronation = request.args.getlist('coronation')
    activity_performed = request.args.getlist('activity_performed')
    observation_category = request.args.getlist('observation_category')

    all_photos = database.get_photos_for_indexing_with_filters(has_coords=True)

    field_titles = {
        'work_front': 'Frente de Trabajo',
        'coronation': 'Coronamiento',
        'activity_performed': 'Actividad Realizada',
        'observation_category': 'Categoría de Observación',
    }
    filter_selections = {
        'work_front': set(work_front),
        'coronation': set(coronation),
        'activity_performed': set(activity_performed),
        'observation_category': set(observation_category),
    }
    date_from = parse_date(date_from_str)
    date_to = parse_date(date_to_str)

    filtered = []
    for photo in all_photos:
        if not photo.get('coords'):
            continue
        pd = photo.get('date')
        if pd and isinstance(pd, str):
            try:
                pd = datetime.fromisoformat(pd)
            except (ValueError, TypeError):
                continue
        if not pd:
            continue
        if date_from and date_to:
            if pd.date() < date_from or pd.date() > date_to:
                continue
            h = pd.hour
            if shift == 'Diurno' and not (7 <= h < 19):
                continue
            if shift == 'Nocturno' and not (h >= 19 or h < 7):
                continue
        elif shift == 'Diurno' and not (7 <= pd.hour < 19):
            continue
        elif shift == 'Nocturno' and not (pd.hour >= 19 or pd.hour < 7):
            continue
        skip = False
        for field, selected in filter_selections.items():
            if not selected:
                continue
            val = photo.get(field)
            norm_val = normalize_metadata_value(val, field_titles[field])
            if norm_val not in selected:
                skip = True
                break
        if skip:
            continue
        filtered.append(photo)

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte de Fotos"

    headers = ["Nombre", "Ruta", "Fecha", "Latitud", "Longitud",
               "Frente de Trabajo", "Coronamiento", "Actividad Realizada",
               "Categoría de Observación"]

    title_font = Font(size=14, bold=True)
    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="E6E6FA", end_color="E6E6FA", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")

    ws.cell(row=1, column=1, value="REPORTE DE FOTOS").font = title_font

    row = 3
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = center_alignment
    row += 1

    for photo in filtered:
        path = photo.get('path', '')
        filename = os.path.basename(path)
        date_val = photo.get('date')
        date_str = date_val.strftime('%Y-%m-%d %H:%M:%S') if date_val else ''
        lat = photo['coords'][0] if photo.get('coords') else ''
        lon = photo['coords'][1] if photo.get('coords') else ''
        wf = photo.get('work_front', '')
        cor = photo.get('coronation', '')
        act = photo.get('activity_performed', '')
        cat = photo.get('observation_category', '')

        row_data = [filename, path, date_str, f"{lat:.6f}" if lat else '', f"{lon:.6f}" if lon else '', wf, cor, act, cat]
        for col, val in enumerate(row_data, 1):
            ws.cell(row=row, column=col, value=val)
        row += 1

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 22

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='reporte_fotos.xlsx')


@app.route('/api/export/table')
def export_table():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    config = load_config()

    if not database.db_exists():
        return jsonify({'error': 'Base de datos no disponible'}), 400

    date_from_str = request.args.get('date_from')
    date_to_str = request.args.get('date_to')
    shift = request.args.get('shift', 'Ambos')
    work_front = request.args.getlist('work_front')
    coronation = request.args.getlist('coronation')
    activity_performed = request.args.getlist('activity_performed')
    observation_category = request.args.getlist('observation_category')

    all_photos = database.get_photos_for_indexing_with_filters(has_coords=True)

    field_titles = {
        'work_front': 'Frente de Trabajo',
        'coronation': 'Coronamiento',
        'activity_performed': 'Actividad Realizada',
        'observation_category': 'Categoría de Observación',
    }
    filter_selections = {
        'work_front': set(work_front),
        'coronation': set(coronation),
        'activity_performed': set(activity_performed),
        'observation_category': set(observation_category),
    }
    date_from = parse_date(date_from_str)
    date_to = parse_date(date_to_str)

    filtered = []
    for photo in all_photos:
        if not photo.get('coords'):
            continue
        pd = photo.get('date')
        if pd and isinstance(pd, str):
            try:
                pd = datetime.fromisoformat(pd)
            except (ValueError, TypeError):
                continue
        if not pd:
            continue
        if date_from and date_to:
            if pd.date() < date_from or pd.date() > date_to:
                continue
            h = pd.hour
            if shift == 'Diurno' and not (7 <= h < 19):
                continue
            if shift == 'Nocturno' and not (h >= 19 or h < 7):
                continue
        elif shift == 'Diurno' and not (7 <= pd.hour < 19):
            continue
        elif shift == 'Nocturno' and not (pd.hour >= 19 or pd.hour < 7):
            continue
        skip = False
        for field, selected in filter_selections.items():
            if not selected:
                continue
            val = photo.get(field)
            norm_val = normalize_metadata_value(val, field_titles[field])
            if norm_val not in selected:
                skip = True
                break
        if skip:
            continue
        filtered.append(photo)

    categories = defaultdict(lambda: defaultdict(list))
    work_fronts_set = set()
    dates_set = set()

    for photo in filtered:
        cat = photo.get('observation_category', '') or 'Sin Categoría'
        wf = photo.get('work_front', '') or 'Sin Frente'
        dt = photo.get('date')
        if dt and isinstance(dt, str):
            try:
                dt = datetime.fromisoformat(dt)
            except (ValueError, TypeError):
                continue
        if not dt:
            continue
        if dt.hour < 7:
            date = (dt - timedelta(days=1)).date()
        else:
            date = dt.date()
        categories[cat][(wf, date)].append(photo)
        work_fronts_set.add(wf)
        dates_set.add(date)

    work_fronts_list = sorted(list(work_fronts_set))
    dates_list = sorted(list(dates_set))

    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True)
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    center_alignment = Alignment(horizontal="center", vertical="center")
    purple_fill = PatternFill(start_color="800080", end_color="800080", fill_type="solid")
    white_font = Font(color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for category, data in categories.items():
        safe_name = (category or "Sin Nombre")[:31]
        ws = wb.create_sheet(title=safe_name)

        ws.cell(row=1, column=1, value="Frente de Trabajo").font = header_font
        ws.cell(row=1, column=1).fill = header_fill
        ws.cell(row=1, column=1).alignment = center_alignment
        ws.cell(row=1, column=1).border = thin_border

        for col_idx, date in enumerate(dates_list, 2):
            c = ws.cell(row=1, column=col_idx, value=date.strftime('%Y-%m-%d'))
            c.font = header_font
            c.fill = header_fill
            c.alignment = center_alignment
            c.border = thin_border

        for row_idx, wf in enumerate(work_fronts_list, 2):
            ws.cell(row=row_idx, column=1, value=wf).font = header_font
            ws.cell(row=row_idx, column=1).fill = header_fill
            ws.cell(row=row_idx, column=1).alignment = center_alignment
            ws.cell(row=row_idx, column=1).border = thin_border

            for col_idx, date in enumerate(dates_list, 2):
                key = (wf, date)
                cell = ws.cell(row=row_idx, column=col_idx)
                if key in data:
                    photos = data[key]
                    act = photos[0].get('activity_performed', '')
                    cor_val = photos[0].get('coronation', '')
                    combined = f"{act} {cor_val}".strip()
                    if combined:
                        cell.value = combined
                        cell.fill = purple_fill
                        cell.font = white_font
                    else:
                        cell.value = ""
                    cell.alignment = center_alignment
                cell.border = thin_border

        for col in range(1, len(dates_list) + 2):
            ws.column_dimensions[get_column_letter(col)].width = 25
        ws.freeze_panes = 'B2'

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name='tabla_actividades.xlsx')


# ── User Management API ──

@app.route('/api/users', methods=['GET'])
def get_users():
    database.init_db()
    database.init_users_table()
    users = database.list_users()
    return jsonify([{'username': u[0], 'created_at': u[1].isoformat() if u[1] else ''} for u in users])


@app.route('/api/users', methods=['POST'])
def create_user_api():
    data = request.get_json()
    username = data.get('username', '').strip()
    password = data.get('password', '')
    if not username or not password:
        return jsonify({'error': 'Usuario y contraseña son obligatorios'}), 400
    if len(username) < 2:
        return jsonify({'error': 'El usuario debe tener al menos 2 caracteres'}), 400
    if len(password) < 4:
        return jsonify({'error': 'La contraseña debe tener al menos 4 caracteres'}), 400
    database.init_db()
    database.init_users_table()
    if not database.create_user(username, password):
        return jsonify({'error': f'El usuario "{username}" ya existe'}), 409
    return jsonify({'status': 'ok', 'username': username})


@app.route('/api/users/<username>', methods=['PUT'])
def update_user_api(username):
    data = request.get_json()
    new_password = data.get('new_password', '')
    if not new_password or len(new_password) < 4:
        return jsonify({'error': 'La nueva contraseña debe tener al menos 4 caracteres'}), 400
    import bcrypt
    password_hash = bcrypt.hashpw(new_password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')
    conn = database.get_connection()
    cursor = conn.cursor()
    cursor.execute(f"SET search_path TO {database.DB_SCHEMA}")
    cursor.execute("UPDATE users SET password_hash = %s WHERE username = %s", (password_hash, username))
    conn.commit()
    affected = cursor.rowcount
    conn.close()
    if affected == 0:
        return jsonify({'error': f'Usuario "{username}" no encontrado'}), 404
    return jsonify({'status': 'ok'})


@app.route('/api/users/<username>', methods=['DELETE'])
def delete_user_api(username):
    current_user = session.get('user', '')
    if username == current_user:
        return jsonify({'error': 'No puede eliminar su propio usuario'}), 403
    database.init_db()
    database.init_users_table()
    database.delete_user(username)
    return jsonify({'status': 'ok'})


# ── Stats ──

@app.route('/api/stats')
def stats():
    try:
        if not database.db_exists():
            return jsonify({'total': 0, 'with_coords': 0})
        database.init_db()
        return jsonify({
            'total': database.get_photo_count(),
            'with_coords': database.get_photos_with_coords_count(),
        })
    except Exception:
        return jsonify({'total': 0, 'with_coords': 0})


if __name__ == '__main__':
    if len(sys.argv) >= 3 and sys.argv[1] == '--add-user':
        username = sys.argv[2]
        password = sys.argv[3] if len(sys.argv) >= 4 else ''
        if not password:
            password = input('Contraseña: ')
        database.init_db()
        database.init_users_table()
        if database.create_user(username, password):
            print(f'Usuario "{username}" creado exitosamente.')
        else:
            print(f'Error: El usuario "{username}" ya existe.')
        sys.exit(0)

    if len(sys.argv) >= 3 and sys.argv[1] == '--delete-user':
        username = sys.argv[2]
        database.init_db()
        database.init_users_table()
        database.delete_user(username)
        print(f'Usuario "{username}" eliminado.')
        sys.exit(0)

    if len(sys.argv) >= 2 and sys.argv[1] == '--list-users':
        database.init_db()
        database.init_users_table()
        users = database.list_users()
        if not users:
            print('No hay usuarios registrados.')
        else:
            for u in users:
                print(f'  {u[0]} (creado: {u[1]})')
        sys.exit(0)

    # Initialize users table on startup
    try:
        database.init_db()
        database.init_users_table()
    except Exception:
        pass

    # Load or generate a persistent secret key
    secret_path = database.get_app_data_path('secret_key.txt')
    if os.path.exists(secret_path):
        with open(secret_path, 'r') as f:
            app.secret_key = f.read().strip()
    else:
        app.secret_key = uuid.uuid4().hex
        with open(secret_path, 'w') as f:
            f.write(app.secret_key)

    app.permanent_session_lifetime = timedelta(hours=8)

    # Use waitress for production, flask dev server as fallback
    try:
        from waitress import serve  # type: ignore[import-untyped]
        print('Iniciando servidor Waitress en http://0.0.0.0:5000')
        print('Presione Ctrl+C para detener.')
        import threading
        import webbrowser
        threading.Timer(1.0, lambda: webbrowser.open('http://localhost:5000')).start()
        serve(app, host='0.0.0.0', port=5000)
    except ImportError:
        print('Waitress no instalado, usando servidor Flask de desarrollo.')
        app.run(debug=True, host='0.0.0.0', port=5000)